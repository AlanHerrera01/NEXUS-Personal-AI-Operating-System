from __future__ import annotations

import docx
import openpyxl
import pytest

from app.config.settings import Settings
from app.domain.ports.tool import ToolContext
from app.domain.value_objects.entity_id import EntityId
from app.infrastructure.skills.files.skill import FilesSkill


def _settings(root) -> Settings:
    return Settings(_env_file=None, files_workspace_root=str(root))


def _tools(settings: Settings) -> dict[str, object]:
    return {tool.definition().name: tool for tool in FilesSkill(settings).tools()}


def _context(user_id: EntityId | None = None) -> ToolContext:
    return ToolContext(
        agent_id=EntityId.new(),
        agent_run_id=EntityId.new(),
        user_id=user_id,
    )


def test_files_skill_exposes_three_read_only_tools(tmp_path) -> None:
    tools = _tools(_settings(tmp_path))

    assert set(tools) == {"files.search", "files.read", "docs.extract"}
    for tool in tools.values():
        definition = tool.definition()
        assert definition.read_only is True
        assert definition.side_effect is False
        assert definition.requires_confirmation is False
        assert definition.input_schema["additionalProperties"] is False


async def test_search_matches_content_and_never_lists_credentials(tmp_path) -> None:
    (tmp_path / "notes.md").write_text(
        "La contraseña del laboratorio es needle-123.\n", encoding="utf-8"
    )
    (tmp_path / ".env").write_text("TOPSECRET=do-not-read\n", encoding="utf-8")
    (tmp_path / "id_rsa.pem").write_text("TOPSECRET\n", encoding="utf-8")
    tools = _tools(_settings(tmp_path))

    result = await tools["files.search"].execute(
        {"query": "needle", "include_content": True}, _context()
    )

    assert result.success is True
    assert result.data["count"] == 1
    assert result.data["items"][0]["path"] == "notes.md"
    assert "needle-123" in result.data["items"][0]["snippet"]

    secret = await tools["files.search"].execute(
        {"query": "TOPSECRET", "include_content": True}, _context()
    )
    assert secret.success is True
    assert secret.data["count"] == 0


async def test_read_rejects_traversal_credentials_and_documents(tmp_path) -> None:
    (tmp_path / "notes.md").write_text("hola NEXUS", encoding="utf-8")
    (tmp_path / ".env").write_text("TOPSECRET=x", encoding="utf-8")
    (tmp_path / "manual.pdf").write_bytes(b"%PDF-1.4")
    tools = _tools(_settings(tmp_path))

    ok = await tools["files.read"].execute({"path": "notes.md"}, _context())
    assert ok.success is True
    assert "hola NEXUS" in ok.data["content"]

    traversal = await tools["files.read"].execute(
        {"path": "../outside.txt"}, _context()
    )
    assert traversal.success is False
    assert traversal.error_code == "FILE_ACCESS_DENIED"

    credentials = await tools["files.read"].execute({"path": ".env"}, _context())
    assert credentials.success is False
    assert credentials.error_code == "FILE_ACCESS_DENIED"

    document = await tools["files.read"].execute({"path": "manual.pdf"}, _context())
    assert document.success is False
    assert document.error_code == "FILE_UNSUPPORTED"


async def test_docs_extract_handles_text_docx_and_xlsx(tmp_path) -> None:
    (tmp_path / "nota.txt").write_text("texto plano", encoding="utf-8")

    document = docx.Document()
    document.add_paragraph("contenido del docx")
    document.save(tmp_path / "informe.docx")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["fila", "xlsx"])
    workbook.save(tmp_path / "planilla.xlsx")

    tools = _tools(_settings(tmp_path))

    text_result = await tools["docs.extract"].execute(
        {"path": "nota.txt"}, _context()
    )
    assert text_result.success is True
    assert text_result.data["format"] == "text"
    assert "texto plano" in text_result.data["content"]

    docx_result = await tools["docs.extract"].execute(
        {"path": "informe.docx"}, _context()
    )
    assert docx_result.success is True
    assert docx_result.data["format"] == "docx"
    assert "contenido del docx" in docx_result.data["content"]

    xlsx_result = await tools["docs.extract"].execute(
        {"path": "planilla.xlsx"}, _context()
    )
    assert xlsx_result.success is True
    assert xlsx_result.data["format"] == "xlsx"
    assert "fila" in xlsx_result.data["content"]


async def test_missing_workspace_is_a_clear_call_time_error(tmp_path) -> None:
    missing = tmp_path / "does-not-exist"
    tools = _tools(_settings(missing))

    result = await tools["files.read"].execute({"path": "notes.md"}, _context())

    assert result.success is False
    assert result.error_code == "FILE_NOT_FOUND"
