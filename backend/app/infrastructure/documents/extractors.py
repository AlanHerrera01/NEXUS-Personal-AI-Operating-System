"""Text extraction for the formats a personal assistant actually receives.

PDF, DOCX and XLSX support comes from pure-Python/desktop libraries. Each
reader is bounded by the caller's byte and character ceilings, and failures are
reported as a domain error rather than a traceback or a half-parsed string.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from app.infrastructure.documents.workspace import (
    TEXT_EXTENSIONS,
    WorkspaceFileTooLarge,
    WorkspaceUnsupported,
)


@dataclass(frozen=True, slots=True)
class TextReadResult:
    text: str
    truncated: bool
    size: int
    modified_at: datetime


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    text: str
    format: str
    truncated: bool
    size: int
    modified_at: datetime
    metadata: dict[str, Any] = field(default_factory=dict)


def read_text_file(
    path: Path,
    *,
    max_bytes: int,
    max_chars: int,
) -> TextReadResult:
    """Read a text file with a byte ceiling and a UTF-8-first decoder."""
    stat = path.stat()
    if stat.st_size > max_bytes:
        raise WorkspaceFileTooLarge(
            f"file is larger than the {max_bytes} byte limit: {path.name}"
        )
    raw = path.read_bytes()
    text = decode_text(raw)
    truncated = len(text) > max_chars
    return TextReadResult(
        text=text[:max_chars],
        truncated=truncated,
        size=stat.st_size,
        modified_at=datetime.fromtimestamp(stat.st_mtime, tz=UTC),
    )


def extract_document(
    path: Path,
    *,
    max_bytes: int,
    max_chars: int,
) -> ExtractionResult:
    """Extract readable text from PDF, DOCX, XLSX or a text file."""
    stat = path.stat()
    if stat.st_size > max_bytes:
        raise WorkspaceFileTooLarge(
            f"file is larger than the {max_bytes} byte limit: {path.name}"
        )
    suffix = path.suffix.lower()
    modified_at = datetime.fromtimestamp(stat.st_mtime, tz=UTC)

    if suffix == ".pdf":
        text, metadata = _extract_pdf(path, max_chars=max_chars)
        format_name = "pdf"
    elif suffix == ".docx":
        text, metadata = _extract_docx(path, max_chars=max_chars)
        format_name = "docx"
    elif suffix == ".xlsx":
        text, metadata = _extract_xlsx(path, max_chars=max_chars)
        format_name = "xlsx"
    elif suffix in TEXT_EXTENSIONS or not suffix:
        text = decode_text(path.read_bytes())
        metadata = {}
        format_name = "text"
    else:
        raise WorkspaceUnsupported(
            f"unsupported document type: {suffix or '(no extension)'}"
        )

    truncated = len(text) > max_chars
    return ExtractionResult(
        text=text[:max_chars],
        format=format_name,
        truncated=truncated,
        size=stat.st_size,
        modified_at=modified_at,
        metadata=metadata,
    )


def decode_text(raw: bytes) -> str:
    if b"\x00" in raw[:2_048]:
        raise WorkspaceUnsupported("file looks binary, not text")
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_pdf(path: Path, *, max_chars: int) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover - dependency is pinned
        raise WorkspaceUnsupported("PDF extraction requires pypdf") from error

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise WorkspaceUnsupported("encrypted PDFs are not supported")
        pages: list[str] = []
        consumed = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            pages.append(page_text)
            consumed += len(page_text)
            if consumed >= max_chars * 2:
                break
        text = "\n\n".join(pages).strip()
    except WorkspaceUnsupported:
        raise
    except Exception as error:
        raise WorkspaceUnsupported("the PDF could not be read") from error

    return text, {"pages": len(reader.pages), "pages_read": len(pages)}


def _extract_docx(path: Path, *, max_chars: int) -> tuple[str, dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover - dependency is pinned
        raise WorkspaceUnsupported("DOCX extraction requires python-docx") from error

    try:
        document = Document(str(path))
        lines: list[str] = []
        consumed = 0
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            lines.append(value)
            consumed += len(value) + 1
            if consumed >= max_chars * 2:
                break
        for table in document.tables:
            if consumed >= max_chars * 2:
                break
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                value = " | ".join(cells).strip()
                if value:
                    lines.append(value)
                    consumed += len(value) + 1
        text = "\n".join(lines).strip()
    except Exception as error:
        raise WorkspaceUnsupported("the DOCX could not be read") from error

    return text, {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
    }


def _extract_xlsx(path: Path, *, max_chars: int) -> tuple[str, dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - dependency is pinned
        raise WorkspaceUnsupported("XLSX extraction requires openpyxl") from error

    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        lines: list[str] = []
        consumed = 0
        rows_read = 0
        for sheet in workbook.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None]
                if not values:
                    continue
                value = "\t".join(values)
                lines.append(value)
                consumed += len(value) + 1
                rows_read += 1
                if consumed >= max_chars * 2:
                    break
            if consumed >= max_chars * 2:
                break
        text = "\n".join(lines).strip()
        sheet_count = len(workbook.sheetnames)
        workbook.close()
    except Exception as error:
        raise WorkspaceUnsupported("the XLSX could not be read") from error

    return text, {"sheets": sheet_count, "rows_read": rows_read}
